from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "docs" / "TinyControlBoard-Documentation.docx"

SOURCE_ORDER = [
    Path("docs/project-guide.md"),
    Path("docs/architecture.md"),
    Path("docs/wiring-reference.md"),
    Path("docs/protocol-reference.md"),
    Path("docs/power-sequencing.md"),
    Path("docs/raspberry-pi-setup.md"),
    Path("scripts/rpi/README.md"),
    Path("docs/button-command-map.md"),
    Path("docs/add-button-how-to.md"),
    Path("docs/project-structure-plan.md"),
]

ACCENT_BLUE = RGBColor(31, 78, 121)
HEADER_FILL = "1F4E79"
ROW_FILL = "EAF2F8"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a single DOCX file from the TinyControlBoard Markdown docs."
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Destination DOCX path.",
    )
    return parser.parse_args()


def ensure_custom_styles(document: Document) -> None:
    styles = document.styles

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code_style.base_style = styles["No Spacing"]
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(9)
        code_style.paragraph_format.left_indent = Inches(0.25)
        code_style.paragraph_format.space_before = Pt(6)
        code_style.paragraph_format.space_after = Pt(6)

    if "Source Note" not in styles:
        source_note = styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
        source_note.base_style = styles["No Spacing"]
        source_note.font.name = "Calibri"
        source_note.font.size = Pt(9)
        source_note.font.italic = True
        source_note.font.color.rgb = RGBColor(96, 96, 96)
        source_note.paragraph_format.space_after = Pt(10)

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for level, size in ((1, 20), (2, 16), (3, 14), (4, 12)):
        heading = styles[f"Heading {level}"]
        heading.font.name = "Calibri"
        heading.font.size = Pt(size)
        heading.font.color.rgb = ACCENT_BLUE


def configure_section(section) -> None:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)


def set_page_layout(section, orientation: WD_ORIENT) -> None:
    if section.orientation == orientation:
        return

    width = section.page_width
    height = section.page_height
    section.orientation = orientation
    section.page_width = height
    section.page_height = width


def set_paragraph_bottom_border(paragraph, color: str = HEADER_FILL) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)

    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)


def content_width_inches(section) -> float:
    usable_width = section.page_width - section.left_margin - section.right_margin
    return usable_width / 914400


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field_run(paragraph, instruction: str, display_text: str = "") -> None:
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)

    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instr_run._r.append(instr)

    separate_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)

    if display_text:
        paragraph.add_run(display_text)

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_page_number(paragraph) -> None:
    paragraph.add_run("Page ")
    add_field_run(paragraph, "PAGE", "1")


def add_title_page(document: Document, sources: list[Path]) -> None:
    banner = document.add_paragraph()
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    banner_run = banner.add_run("PROJECT DOCUMENTATION PACK")
    banner_run.bold = True
    banner_run.font.size = Pt(10)
    banner_run.font.color.rgb = ACCENT_BLUE

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TinyControlBoard Documentation")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = ACCENT_BLUE

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Consolidated engineering reference and implementation guide")

    generated = document.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated_run = generated.add_run(f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    generated_run.italic = True

    document.add_paragraph()
    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run("Included source documents")

    for source in sources:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(str(source).replace("\\", "/"))
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(96, 96, 96)

    document.add_page_break()
    heading = document.add_paragraph(style="Heading 1")
    heading.add_run("Table Of Contents")
    toc = document.add_paragraph()
    add_field_run(toc, 'TOC \\o "1-3" \\h \\z \\u', "Update fields in Word to populate the table of contents.")
    document.add_page_break()


def add_inline_content(paragraph, node, base_dir: Path) -> None:
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            paragraph.add_run(text)
        return

    if not isinstance(node, Tag):
        return

    if node.name == "br":
        paragraph.add_run().add_break()
        return

    if node.name == "code":
        run = paragraph.add_run(node.get_text())
        run.font.name = "Consolas"
        return

    if node.name == "strong":
        start = len(paragraph.runs)
        for child in node.children:
            add_inline_content(paragraph, child, base_dir)
        for run in paragraph.runs[start:]:
            run.bold = True
        return

    if node.name == "em":
        start = len(paragraph.runs)
        for child in node.children:
            add_inline_content(paragraph, child, base_dir)
        for run in paragraph.runs[start:]:
            run.italic = True
        return

    if node.name == "a":
        label = node.get_text(strip=False) or node.get("href", "")
        run = paragraph.add_run(label)
        run.underline = True
        return

    for child in node.children:
        add_inline_content(paragraph, child, base_dir)


def add_paragraph_from_tag(document: Document, tag: Tag, base_dir: Path, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    for child in tag.children:
        add_inline_content(paragraph, child, base_dir)
    return paragraph


def add_source_note(document: Document, source: Path) -> None:
    paragraph = document.add_paragraph(style="Source Note")
    paragraph.add_run(f"Source: {source.relative_to(ROOT).as_posix()}")


def add_list(document: Document, list_tag: Tag, base_dir: Path, ordered: bool, level: int = 0) -> None:
    style = "List Number" if ordered else "List Bullet"
    for item in list_tag.find_all("li", recursive=False):
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.left_indent = Inches(0.3 * level)
        nested_lists: list[Tag] = []
        for child in item.children:
            if isinstance(child, Tag) and child.name in {"ul", "ol"}:
                nested_lists.append(child)
                continue
            add_inline_content(paragraph, child, base_dir)

        for nested in nested_lists:
            add_list(document, nested, base_dir, nested.name == "ol", level + 1)


def add_code_block(document: Document, code_tag: Tag) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    text = code_tag.get_text().rstrip("\n")
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    shade_paragraph(paragraph, "F4F6F7")


def add_image(document: Document, image_tag: Tag, base_dir: Path) -> None:
    src = image_tag.get("src")
    if not src:
        return

    image_path = (base_dir / src).resolve()
    if not image_path.exists():
        paragraph = document.add_paragraph()
        paragraph.add_run(f"[Missing image: {src}]")
        return

    width = min(content_width_inches(document.sections[-1]), 6.8)
    document.add_picture(str(image_path), width=Inches(width))
    if image_tag.get("alt"):
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(image_tag["alt"]).italic = True


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_table(document: Document, table_tag: Tag, base_dir: Path) -> None:
    landscape = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(landscape)
    set_page_layout(landscape, WD_ORIENT.LANDSCAPE)

    rows = table_tag.find_all("tr")
    if not rows:
        portrait = document.add_section(WD_SECTION.NEW_PAGE)
        set_page_layout(portrait, WD_ORIENT.PORTRAIT)
        return

    max_cols = max(len(row.find_all(["th", "td"], recursive=False)) for row in rows)
    table = document.add_table(rows=0, cols=max_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    available_width = (
        landscape.page_width - landscape.left_margin - landscape.right_margin
    )
    column_char_counts = [1] * max_cols
    extracted_rows: list[list[Tag]] = []

    for row in rows:
        cells = row.find_all(["th", "td"], recursive=False)
        extracted_rows.append(cells)
        for index, cell in enumerate(cells):
            column_char_counts[index] = max(
                column_char_counts[index], len(cell.get_text(" ", strip=True))
            )

    total_chars = sum(column_char_counts)
    column_widths = [int(available_width * count / total_chars) for count in column_char_counts]
    remainder = int(available_width) - sum(column_widths)
    if column_widths:
        column_widths[-1] += remainder

    for row_index, html_cells in enumerate(extracted_rows):
        row = table.add_row()
        for col_index in range(max_cols):
            cell = row.cells[col_index]
            cell.width = column_widths[col_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["No Spacing"]
            if col_index < len(html_cells):
                cell_tag = html_cells[col_index]
                for child in cell_tag.children:
                    add_inline_content(paragraph, child, base_dir)

            for run in paragraph.runs:
                run.font.size = Pt(9)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)

            if row_index == 0:
                set_cell_shading(cell, HEADER_FILL)
            elif row_index % 2 == 1:
                set_cell_shading(cell, ROW_FILL)

    if table.rows:
        set_repeat_table_header(table.rows[0])

    portrait = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(portrait)
    set_page_layout(portrait, WD_ORIENT.PORTRAIT)


def render_html_block(document: Document, node, base_dir: Path) -> None:
    if isinstance(node, NavigableString):
        if str(node).strip():
            paragraph = document.add_paragraph()
            paragraph.add_run(str(node).strip())
        return

    if not isinstance(node, Tag):
        return

    if node.name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        level = min(int(node.name[1]), 4)
        add_paragraph_from_tag(document, node, base_dir, style=f"Heading {level}")
        return

    if node.name == "p":
        if node.find("img") and len(node.contents) == 1:
            add_image(document, node.img, base_dir)
            return
        add_paragraph_from_tag(document, node, base_dir)
        return

    if node.name == "ul":
        add_list(document, node, base_dir, ordered=False)
        return

    if node.name == "ol":
        add_list(document, node, base_dir, ordered=True)
        return

    if node.name == "pre":
        code = node.find("code") or node
        add_code_block(document, code)
        return

    if node.name == "table":
        add_table(document, node, base_dir)
        return

    if node.name == "hr":
        document.add_paragraph()
        return

    for child in node.children:
        render_html_block(document, child, base_dir)


def render_markdown_file(document: Document, source: Path) -> None:
    markdown_text = source.read_text(encoding="utf-8")
    html = markdown.markdown(
        markdown_text,
        extensions=["tables", "fenced_code", "sane_lists"],
        output_format="html5",
    )
    soup = BeautifulSoup(html, "html.parser")

    add_source_note(document, source)

    for node in soup.contents:
        render_html_block(document, node, source.parent)


def add_headers_and_footers(document: Document) -> None:
    for section in document.sections:
        configure_section(section)

        header = section.header
        if header.paragraphs:
            header_paragraph = header.paragraphs[0]
            header_paragraph.clear()
        else:
            header_paragraph = header.add_paragraph()
        header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_run = header_paragraph.add_run("TinyControlBoard Documentation")
        header_run.bold = True
        header_run.font.size = Pt(9)
        header_run.font.color.rgb = ACCENT_BLUE
        set_paragraph_bottom_border(header_paragraph)

        footer = section.footer
        if footer.paragraphs:
            footer_paragraph = footer.paragraphs[0]
            footer_paragraph.clear()
        else:
            footer_paragraph = footer.add_paragraph()
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer_paragraph.add_run("TinyControlBoard")
        footer_run.font.size = Pt(9)
        footer_run.font.color.rgb = RGBColor(96, 96, 96)
        footer_paragraph.add_run("  |  ")
        add_page_number(footer_paragraph)


def build_document(output_path: Path) -> Path:
    sources = [ROOT / relative_path for relative_path in SOURCE_ORDER if (ROOT / relative_path).exists()]
    document = Document()
    ensure_custom_styles(document)
    document.core_properties.title = "TinyControlBoard Documentation"
    document.core_properties.subject = "Combined Markdown documentation export"
    document.core_properties.author = "GitHub Copilot"

    first_section = document.sections[0]
    configure_section(first_section)

    add_title_page(document, [source.relative_to(ROOT) for source in sources])

    for index, source in enumerate(sources):
        if index > 0:
            document.add_page_break()
        render_markdown_file(document, source)

    add_headers_and_footers(document)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def main() -> int:
    args = parse_args()
    output_path = args.output
    if not output_path.is_absolute():
        output_path = (ROOT / output_path).resolve()

    saved_path = build_document(output_path)
    print(f"Wrote {saved_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())