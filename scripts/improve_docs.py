"""
improve_docs.py - Makes TinyControlBoard-Documentation.docx more professional.

Changes applied:
  1. Cover page: proper title block replacing auto-generated text.
  2. Fix "setup Moode (Raspberry Pi)" Heading1 to proper Appendix title.
  3. Fix the unnumbered/lowercase "add meters" heading.
  4. Fix per-button notes: inline dash-separated text to proper bullet lists.
  5. Fix typos in the Pi setup section.
  6. Remove "PROJECT DOCUMENTATION PACK" / "Generated..." metadata text.
  7. Apply Heading4 style to troubleshooting sub-headings.
  8. Fix "Repo Layout" heading.
"""

import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

SRC = r"d:\Dev\TinyControlBoard\docs\TinyControlBoard-Documentation.docx"
DST = r"d:\Dev\TinyControlBoard\docs\TinyControlBoard-Documentation-Professional.docx"

doc = Document(SRC)

def para_text(p):
    return "".join(r.text or "" for r in p.runs)

def set_para_text(p, text):
    """Clear all runs (including embedded line breaks) and set plain text."""
    from docx.oxml.ns import qn as qn_
    for r in p.runs:
        r.text = ""
        # Remove any <w:br/> elements left inside the run
        for br in r._element.findall(qn_("w:br")):
            r._element.remove(br)
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)

def apply_style(p, style_name):
    try:
        p.style = doc.styles[style_name]
    except Exception:
        pass

def delete_para(p):
    p._element.getparent().remove(p._element)

def iter_all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p

def insert_bullets_after(ref_p, items, doc):
    ref_elem = ref_p._element
    for item in reversed(items):
        new_p_elem = OxmlElement("w:p")
        ref_elem.addnext(new_p_elem)
        new_para = Paragraph(new_p_elem, doc)
        new_para.add_run(item)
        apply_style(new_para, "List Bullet")

# 1. Fix cover page
cover_remove = {"PROJECT DOCUMENTATION PACK", "Included source documents", "docs/TinyControlBoard-Documentation.md"}
to_delete = []
for p in doc.paragraphs:
    t = para_text(p).strip()
    if t in cover_remove:
        to_delete.append(p)
    elif t.startswith("Generated "):
        set_para_text(p, "Revision 1.0  -  2026-06-05")
        apply_style(p, "Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for p in to_delete:
    delete_para(p)

for p in doc.paragraphs:
    if para_text(p).strip() == "TinyControlBoard Documentation":
        apply_style(p, "Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        break

for p in doc.paragraphs:
    if para_text(p).strip() == "Consolidated engineering reference and implementation guide":
        apply_style(p, "Subtitle")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        break

# 2. Fix moode heading
for p in iter_all_paragraphs(doc):
    if para_text(p).strip() == "setup Moode (Raspberry Pi)":
        set_para_text(p, "Appendix A -- Moode Audio Setup (Raspberry Pi)")
        break

# 3. Fix add meters heading
for p in iter_all_paragraphs(doc):
    if para_text(p).strip() == "add meters":
        set_para_text(p, "A.14b -- Add Meter Configurations")
        break

# 4. Fix Repo Layout heading
for p in iter_all_paragraphs(doc):
    if para_text(p).strip() == "Repo Layout":
        set_para_text(p, "A.0 -- Repository Layout")
        apply_style(p, "Heading3")
        break

# 5. Fix per-button note paragraphs
INLINE_NOTE_PREFIXES = [
    "Power button (index 0):",
    "Toggle buttons (Cover, Repeat, Random, DAC, Meter):",
    "Rotary (indices 13/14):",
    "SPI LED notes:",
    "Notes:-",
]

def split_note_para(p):
    """Expand 'Header:\n- item1\n- item2' into a Heading4 + bullet list."""
    text = para_text(p).strip()
    # Items are separated by newline-dash: "Header:\n- item1\n- item2"
    newline_idx = text.find("\n-")
    if newline_idx == -1:
        # Fallback: old inline dash format "Header:- item1.- item2"
        colon_pos = text.find(":-")
        if colon_pos == -1:
            return
        header = text[:colon_pos + 1].strip()
        rest = text[colon_pos + 1:].strip()
        items_raw = re.split(r"(?<=[.!?])\s*-\s*", rest)
    else:
        header = text[:newline_idx].strip()
        rest = text[newline_idx:]
        items_raw = rest.split("\n-")
    items = []
    for raw in items_raw:
        it = raw.strip().lstrip("-").strip()
        if it:
            if not it.endswith((".", "!", "?")):
                it += "."
            items.append(it)
    if not items:
        return
    set_para_text(p, header)
    apply_style(p, "Heading4")
    insert_bullets_after(p, items, doc)

note_paras = []
for p in list(iter_all_paragraphs(doc)):
    t = para_text(p).strip()
    for prefix in INLINE_NOTE_PREFIXES:
        if t.startswith(prefix):
            note_paras.append(p)
            break
print(f"Found {len(note_paras)} inline note paragraphs")
for p in note_paras:
    split_note_para(p)

# 6. Fix Actions list paragraphs
def expand_action_list(p, label):
    text = para_text(p).strip()
    body = text[len(label):].strip()
    if not body:
        return
    set_para_text(p, label)
    apply_style(p, "Heading4")
    items = [x.strip().rstrip(".") + "." for x in body.split(",") if x.strip()]
    insert_bullets_after(p, items, doc)

for p in list(iter_all_paragraphs(doc)):
    t = para_text(p).strip()
    if t.startswith("Actions that send UART commands to the Pi:"):
        expand_action_list(p, "Actions that send UART commands to the Pi:")
        break

for p in list(iter_all_paragraphs(doc)):
    t = para_text(p).strip()
    if t.startswith("Actions that stay local on the ESP32:"):
        expand_action_list(p, "Actions that stay local on the ESP32:")
        break

# 7. Typo fixes
TYPO_MAP = [
    ("make sure that the you  fcopy the file from the location where you originaly stored the file",
     "Copy the file from the location where you originally stored it."),
    ("originaly", "originally"),
    ("UAart5Listener.py", "uart5_listener.py"),
    ("code References:", "Code references:"),
    ("code references :", "Code references:"),
    ("/home/antho/UAart5Listener.py", "/home/antho/uart5_listener.py"),
]

# First pass: try per-run replacement (fast path for text within a single run)
for p in iter_all_paragraphs(doc):
    for run in p.runs:
        if run.text:
            t = run.text
            for bad, good in TYPO_MAP:
                t = t.replace(bad, good)
            run.text = t

# Second pass: whole-paragraph replacement for typos split across runs
PARA_TYPO_MAP = [
    ("make sure that the you  fcopy the file from the location where you originaly stored the file",
     "Copy the file from the location where you originally stored it."),
]
for p in iter_all_paragraphs(doc):
    full = para_text(p)
    for bad, good in PARA_TYPO_MAP:
        if bad in full:
            set_para_text(p, full.replace(bad, good))
            break

# 8. Troubleshooting sub-headings
TROUBLESHOOT_HEADS = ["15.1", "15.2", "15.3", "15.4"]
for p in iter_all_paragraphs(doc):
    t = para_text(p).strip()
    if any(t.startswith(h) for h in TROUBLESHOOT_HEADS):
        apply_style(p, "Heading4")

# 9. "Notes:" inline paragraphs to heading
for p in list(doc.paragraphs):
    t = para_text(p).strip()
    if t.startswith("Notes:-") or t.startswith("Notes: -"):
        colon_pos = t.find(":-")
        if colon_pos == -1:
            colon_pos = t.find(": -")
        header = t[:colon_pos + 1].strip()
        rest = t[colon_pos + 1:].strip()
        items_raw = re.split(r"(?<=[.!?])\s*-\s*", rest)
        items = []
        for raw in items_raw:
            item = raw.strip().lstrip("-").strip()
            if item:
                if not item.endswith((".", "!", "?")):
                    item += "."
                items.append(item)
        if items:
            set_para_text(p, header)
            apply_style(p, "Heading4")
            insert_bullets_after(p, items, doc)

doc.save(DST)
print(f"Saved: {DST}")
